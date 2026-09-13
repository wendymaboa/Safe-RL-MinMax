"""
Prepend Wendy_Prior_Updates_styled.pptx in front of a COPY of
Wendy_Progress_Update_Sep2026.pptx via OPC zip merge.

Does NOT modify Wendy_Progress_Update_Sep2026.pptx.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

SLIDES = Path(r'c:\Users\Wendy\Desktop\RLHC\docs\Slides')
SEP = SLIDES / 'Wendy_Progress_Update_Sep2026.pptx'
PRIOR = SLIDES / 'Wendy_Prior_Updates_styled.pptx'
OUT = SLIDES / 'Wendy_Talk_Sep2026_full_trail.pptx'
SEP_BLANK_LAYOUT = '../slideLayouts/slideLayout11.xml'


def read_zip(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def main() -> None:
    if not PRIOR.exists():
        raise SystemExit(f'Missing {PRIOR} — run build_prior_styled.py first')
    if not SEP.exists():
        raise SystemExit(f'Missing {SEP}')

    sep = read_zip(SEP)
    prior = read_zip(PRIOR)

    prior_slides = sorted(
        [n for n in prior if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)],
        key=lambda n: int(re.search(r'(\d+)', n).group(1)),
    )
    sep_slides = sorted(
        [n for n in sep if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)],
        key=lambda n: int(re.search(r'(\d+)', n).group(1)),
    )
    print(f'Prior slides: {len(prior_slides)}, Sep slides: {len(sep_slides)}')

    out: dict[str, bytes] = {
        k: v for k, v in sep.items()
        if 'printerSettings' not in k
        and not re.fullmatch(r'ppt/slides/slide\d+\.xml', k)
        and not re.fullmatch(r'ppt/slides/_rels/slide\d+\.xml.rels', k)
    }

    # Copy prior media with prefix to avoid collisions
    media_overrides: list[str] = []
    for name, data in prior.items():
        if name.startswith('ppt/media/'):
            new_name = f'ppt/media/prior_{Path(name).name}'
            out[new_name] = data
            ext = Path(name).suffix.lower()
            ctype = {
                '.png': 'image/png',
                '.jpg': 'image/jpeg',
                '.jpeg': 'image/jpeg',
                '.gif': 'image/gif',
                '.emf': 'image/x-emf',
                '.wmf': 'image/x-wmf',
                '.svg': 'image/svg+xml',
            }.get(ext)
            if ctype:
                media_overrides.append(
                    f'<Override PartName="/{new_name}" ContentType="{ctype}"/>'
                )

    pres_xml = sep['ppt/presentation.xml'].decode('utf-8')
    pres_xml = pres_xml.replace(
        "<?xml version='1.0' encoding='UTF-8' standalone='yes'?>",
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
    )
    if '<p:sldSz' not in pres_xml:
        insert = '<p:sldSz cx="12191695" cy="6858000"/><p:notesSz cx="6858000" cy="9144000"/>'
        if '<p:defaultTextStyle' in pres_xml:
            pred = '<p:defaultTextStyle'
            pres_xml = pres_xml.replace(pred, insert + pred, 1)

    rels_xml = sep['ppt/_rels/presentation.xml.rels'].decode('utf-8')
    rels_xml = re.sub(r'<Relationship[^>]*Type="[^"]*?/slide"[^>]*/>', '', rels_xml)
    rels_xml = re.sub(r'<Relationship[^>]*printerSettings[^>]*/>', '', rels_xml)

    ct_xml = sep['[Content_Types].xml'].decode('utf-8')
    ct_xml = re.sub(r'<Override[^>]*printerSettings[^>]*/>', '', ct_xml)
    ct_xml = re.sub(r'<Override PartName="/ppt/slides/slide\d+\.xml"[^>]*/>', '', ct_xml)

    def remap_media(rels_text: str) -> str:
        def fix_target(target: str) -> str:
            if '/media/' in target:
                fname = target.split('/')[-1]
                return f'../media/prior_{fname}'
            return target

        return re.sub(
            r'Target="([^"]+)"',
            lambda m: f'Target="{fix_target(m.group(1))}"',
            rels_text,
        )

    slide_rels_entries: list[str] = []
    sld_id_entries: list[str] = []
    ct_overrides: list[str] = []
    next_rid = 100
    next_slide_num = 1
    slide_id_num = 256

    def add_slide(slide_xml: bytes, slide_rels: bytes | None, *, remap_layout: bool) -> None:
        nonlocal next_rid, next_slide_num, slide_id_num
        slide_name = f'ppt/slides/slide{next_slide_num}.xml'
        rels_name = f'ppt/slides/_rels/slide{next_slide_num}.xml.rels'
        out[slide_name] = slide_xml
        if slide_rels is not None:
            text = slide_rels.decode('utf-8')
            if remap_layout:
                text = re.sub(
                    r'(Type="[^"]*?/slideLayout"[^>]*Target=")[^"]+(")',
                    rf'\1{SEP_BLANK_LAYOUT}\2',
                    text,
                )
                text = remap_media(text)
                text = re.sub(r'<Relationship[^>]*notesSlide[^>]*/>', '', text)
            out[rels_name] = text.encode('utf-8')
        else:
            out[rels_name] = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                f'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="{SEP_BLANK_LAYOUT}"/>'
                '</Relationships>'
            ).encode('utf-8')

        rid = f'rId{next_rid}'
        next_rid += 1
        slide_rels_entries.append(
            f'<Relationship Id="{rid}" '
            f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" '
            f'Target="slides/slide{next_slide_num}.xml"/>'
        )
        sld_id_entries.append(f'<p:sldId id="{slide_id_num}" r:id="{rid}"/>')
        ct_overrides.append(
            f'<Override PartName="/{slide_name}" '
            f'ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        )
        slide_id_num += 1
        next_slide_num += 1

    for spath in prior_slides:
        num = re.search(r'(\d+)', Path(spath).name).group(1)
        add_slide(prior[spath], prior.get(f'ppt/slides/_rels/slide{num}.xml.rels'), remap_layout=True)

    for spath in sep_slides:
        num = re.search(r'(\d+)', Path(spath).name).group(1)
        add_slide(sep[spath], sep.get(f'ppt/slides/_rels/slide{num}.xml.rels'), remap_layout=False)

    new_list = '<p:sldIdLst>' + ''.join(sld_id_entries) + '</p:sldIdLst>'
    if not re.search(r'<p:sldIdLst>.*?</p:sldIdLst>', pres_xml, flags=re.S):
        raise SystemExit('Could not find sldIdLst')
    pres_xml = re.sub(r'<p:sldIdLst>.*?</p:sldIdLst>', new_list, pres_xml, flags=re.S)

    rels_xml = rels_xml.replace('</Relationships>', ''.join(slide_rels_entries) + '</Relationships>')
    ct_xml = ct_xml.replace('</Types>', ''.join(media_overrides) + ''.join(ct_overrides) + '</Types>')

    out['ppt/presentation.xml'] = pres_xml.encode('utf-8')
    out['ppt/_rels/presentation.xml.rels'] = rels_xml.encode('utf-8')
    out['[Content_Types].xml'] = ct_xml.encode('utf-8')

    if OUT.exists():
        OUT.unlink()
    with zipfile.ZipFile(OUT, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in out.items():
            zout.writestr(name, data)

    # Verify with python-pptx
    from pptx import Presentation
    prs = Presentation(str(OUT))
    print(f'Wrote {OUT.name}: python-pptx sees {len(prs.slides)} slides')
    print(f'Original untouched: {SEP.name} ({SEP.stat().st_size} bytes)')
    print('Excluded: White Beige TEAM A template (not research content).')


if __name__ == '__main__':
    main()
